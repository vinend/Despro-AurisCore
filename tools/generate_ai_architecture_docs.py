from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "ai_architecture"
DIAGRAM_DIR = OUT_DIR / "diagrams"

TITLE = "Arsitektur AI Stetoskop Digital Multi-Mode"
SUBTITLE = "Heart, Lung, Abdomen Sounds | PDS 2 Smartphone Inference dan PDS 3 Edge AI / TinyML"


PALETTE = {
    "bg": "#ffffff",
    "canvas": "#fbfcfd",
    "grid": "#eef2f6",
    "ink": "#1f2933",
    "muted": "#5f6b7a",
    "line": "#52616f",
    "box": "#ffffff",
    "box_border": "#c5ced8",
    "shadow": "#dce3ea",
    "blue": "#e9f4ff",
    "blue_dark": "#2166a5",
    "green": "#edf8ef",
    "green_dark": "#2f7d4f",
    "orange": "#fff4e5",
    "orange_dark": "#b66813",
    "red": "#fff0f0",
    "red_dark": "#b34545",
    "gray": "#f3f5f7",
    "purple": "#f2f0ff",
    "purple_dark": "#6858b8",
}


def ensure_dirs() -> None:
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf",
        "C:/Windows/Fonts/timesbi.ttf" if bold else "C:/Windows/Fonts/timesi.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


FONT_TITLE = load_font(44, True)
FONT_SUBTITLE = load_font(24, False)
FONT_BOX_TITLE = load_font(25, True)
FONT_BOX = load_font(22, False)
FONT_SMALL = load_font(18, False)
FONT_CAPTION = load_font(19, False)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> tuple[int, int]:
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrapped_lines(text: str, max_chars: int) -> list[str]:
    lines: list[str] = []
    for part in text.split("\n"):
        lines.extend(wrap(part, width=max_chars) if part else [""])
    return lines


def draw_text_center(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
    max_chars: int,
    line_gap: int = 6,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrapped_lines(text, max_chars)
    heights = [text_size(draw, line, font)[1] for line in lines]
    total_h = sum(heights) + line_gap * max(0, len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) // 2
    for line, h in zip(lines, heights):
        w, _ = text_size(draw, line, font)
        draw.text((x1 + ((x2 - x1) - w) // 2, y), line, fill=fill, font=font)
        y += h + line_gap


def draw_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    fill: str,
    outline: str,
    title_color: str | None = None,
) -> None:
    x1, y1, x2, y2 = xy
    shadow = (x1 + 8, y1 + 8, x2 + 8, y2 + 8)
    draw.rounded_rectangle(shadow, radius=18, fill=PALETTE["shadow"])
    draw.rounded_rectangle(xy, radius=18, fill=PALETTE["box"], outline=PALETTE["box_border"], width=2)
    draw.rounded_rectangle((x1, y1, x2, y1 + 18), radius=18, fill=outline)
    draw.rectangle((x1, y1 + 10, x2, y1 + 18), fill=outline)
    draw.rounded_rectangle((x1 + 16, y1 + 28, x1 + 44, y1 + 56), radius=14, fill=fill, outline=outline, width=2)
    title_color = title_color or outline
    title_chars = max(8, (x2 - x1 - 80) // 12)
    body_chars = max(9, (x2 - x1 - 48) // 12)
    draw_text_center(draw, (x1 + 52, y1 + 24, x2 - 16, y1 + 64), title, FONT_BOX_TITLE, title_color, title_chars)
    draw_text_center(draw, (x1 + 24, y1 + 70, x2 - 24, y2 - 18), body, FONT_BOX, PALETTE["ink"], body_chars)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = PALETTE["line"]) -> None:
    sx, sy = start
    ex, ey = end
    if sx == ex or sy == ey:
        points = [start, end]
    else:
        mid_x = sx + (ex - sx) // 2
        points = [start, (mid_x, sy), (mid_x, ey), end]
    draw.line(points, fill=color, width=4, joint="curve")
    if abs(ex - points[-2][0]) >= abs(ey - points[-2][1]):
        direction = 1 if ex >= points[-2][0] else -1
        arrowhead = [(ex, ey), (ex - 17 * direction, ey - 10), (ex - 17 * direction, ey + 10)]
    else:
        direction = 1 if ey >= points[-2][1] else -1
        arrowhead = [(ex, ey), (ex - 10, ey - 17 * direction), (ex + 10, ey - 17 * direction)]
    draw.polygon(arrowhead, fill=color)


def draw_grid(draw: ImageDraw.ImageDraw) -> None:
    for x in range(0, 1800, 60):
        draw.line((x, 0, x, 1050), fill=PALETTE["grid"], width=1)
    for y in range(0, 1050, 60):
        draw.line((0, y, 1800, y), fill=PALETTE["grid"], width=1)


def draw_label(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], text: str, color: str) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=16, fill=color)
    draw_text_center(draw, (x1 + 10, y1, x2 - 10, y2), text, FONT_SMALL, "#ffffff", 30)


def draw_note(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], title: str, body: str) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle((x1 + 6, y1 + 6, x2 + 6, y2 + 6), radius=18, fill=PALETTE["shadow"])
    draw.rounded_rectangle(xy, radius=18, fill="#fffdf7", outline="#d5c49a", width=2)
    title_chars = max(24, (x2 - x1 - 36) // 13)
    body_chars = max(36, (x2 - x1 - 44) // 12)
    draw_text_center(draw, (x1 + 18, y1 + 16, x2 - 18, y1 + 52), title, FONT_BOX_TITLE, PALETTE["ink"], title_chars)
    draw_text_center(draw, (x1 + 22, y1 + 58, x2 - 22, y2 - 16), body, FONT_BOX, PALETTE["muted"], body_chars)


def draw_swimlane(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    label: str,
    color: str,
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill="#ffffff", outline="#d7dde5", width=2)
    draw.rounded_rectangle((x1, y1, x2, y1 + 50), radius=24, fill=color)
    draw.rectangle((x1, y1 + 25, x2, y1 + 50), fill=color)
    draw_text_center(draw, (x1 + 20, y1, x2 - 20, y1 + 50), label, FONT_BOX_TITLE, "#ffffff", 60)


def draw_split_arrow(
    draw: ImageDraw.ImageDraw,
    source: tuple[int, int],
    targets: list[tuple[int, int]],
    color: str = PALETTE["line"],
) -> None:
    sx, sy = source
    hub_x = sx + 70
    draw.line((source, (hub_x, sy)), fill=color, width=4)
    for tx, ty in targets:
        draw.line(((hub_x, sy), (hub_x, ty), (tx, ty)), fill=color, width=4)
        direction = 1 if tx >= hub_x else -1
        draw.polygon([(tx, ty), (tx - 17 * direction, ty - 10), (tx - 17 * direction, ty + 10)], fill=color)


def canvas(title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (1800, 1050), PALETTE["canvas"])
    draw = ImageDraw.Draw(img)
    draw_grid(draw)
    draw.rectangle((0, 0, 1800, 132), fill="#ffffff")
    draw.rectangle((0, 128, 1800, 132), fill="#d6dde6")
    draw.rounded_rectangle((58, 38, 74, 94), radius=8, fill=PALETTE["blue_dark"])
    draw.text((96, 32), title, fill=PALETTE["ink"], font=FONT_TITLE)
    if subtitle:
        draw.text((98, 84), subtitle, fill=PALETTE["muted"], font=FONT_SUBTITLE)
    return img, draw


def save(img: Image.Image, filename: str) -> None:
    img.save(DIAGRAM_DIR / filename, quality=95)


def draw_auscultation_modes() -> None:
    img, draw = canvas("Usage and Auscultation Mode Strategy", "Mode akustik original dan mode digital untuk jantung, paru, dan abdomen")
    draw_note(
        draw,
        (80, 165, 1720, 330),
        "Inspirasi Littmann Digital",
        "Perangkat tetap mempertahankan fungsi dasar stetoskop: dokter dapat mendengar suara tubuh secara langsung. Mode digital ditambahkan untuk perekaman, visualisasi HP/PC, dan AI assistive screening.",
    )
    draw_swimlane(draw, (80, 365, 1720, 660), "Usage Mode", PALETTE["blue_dark"])
    usage_boxes = [
        ((125, 465, 430, 625), "Acoustic Mode", "Auskultasi langsung\nmanual"),
        ((560, 465, 865, 625), "Digital Mode", "Recording +\nvisual HP/PC"),
        ((995, 465, 1300, 625), "PDS 2", "AI di HP/PC\nWebSocket"),
        ((1430, 465, 1680, 625), "PDS 3", "AI di MCU\nhasil ke HP/PC"),
    ]
    usage_fills = [PALETTE["gray"], PALETTE["blue"], PALETTE["green"], PALETTE["orange"]]
    usage_outlines = [PALETTE["line"], PALETTE["blue_dark"], PALETTE["green_dark"], PALETTE["orange_dark"]]
    for (xy, title, body), fill, outline in zip(usage_boxes, usage_fills, usage_outlines):
        draw_box(draw, xy, title, body, fill, outline)
    arrow(draw, (430, 545), (560, 545))
    arrow(draw, (865, 545), (995, 545))
    arrow(draw, (1300, 545), (1430, 545))

    draw_swimlane(draw, (80, 705, 1720, 980), "Body Sound Targets", PALETTE["purple_dark"])
    target_boxes = [
        ((125, 805, 430, 945), "Heart Mode", "S1/S2, murmur\nNormal / abnormal"),
        ((560, 805, 865, 945), "Lung Mode", "Wheeze, crackle\nNormal / abnormal"),
        ((995, 805, 1300, 945), "Abdomen Mode", "Bowel burst\nNormal / abnormal"),
        ((1430, 805, 1680, 945), "Router", "Pilih pipeline\nsesuai lokasi"),
    ]
    target_fills = [PALETTE["red"], PALETTE["blue"], PALETTE["green"], PALETTE["orange"]]
    target_outlines = [PALETTE["red_dark"], PALETTE["blue_dark"], PALETTE["green_dark"], PALETTE["orange_dark"]]
    for (xy, title, body), fill, outline in zip(target_boxes, target_fills, target_outlines):
        draw_box(draw, xy, title, body, fill, outline)
    arrow(draw, (430, 875), (560, 875))
    arrow(draw, (865, 875), (995, 875))
    arrow(draw, (1300, 875), (1430, 875))
    save(img, "00_auscultation_mode_strategy.png")


def draw_overall_architecture() -> None:
    img, draw = canvas("Overall AI Architecture", "Perbandingan jalur PDS 2 dan PDS 3 untuk stetoskop digital multi-mode")

    draw_label(draw, (90, 165, 435, 215), "COMMON ACQUISITION", PALETTE["blue_dark"])
    draw_box(draw, (110, 410, 360, 570), "Body Sounds", "Heart / lung / abdomen sounds", PALETTE["gray"], PALETTE["line"])
    draw_box(draw, (455, 410, 705, 570), "Chestpiece + Mic", "Acoustic coupling + MEMS mic", PALETTE["blue"], PALETTE["blue_dark"])
    draw_box(draw, (455, 650, 705, 800), "Acoustic Mode", "Dokter mendengar\nseperti stetoskop biasa", PALETTE["gray"], PALETTE["line"])
    arrow(draw, (360, 490), (455, 490))
    arrow(draw, (580, 570), (580, 650))

    draw_swimlane(draw, (805, 190, 1690, 465), "PDS 2 - Smartphone / PC Inference", PALETTE["green_dark"])
    draw_box(draw, (850, 285, 1085, 425), "MCU + Wi-Fi", "Capture audio dan stream data", PALETTE["blue"], PALETTE["blue_dark"])
    draw_box(draw, (1165, 285, 1400, 425), "Mode Features", "MFCC / log-mel + scaler", PALETTE["green"], PALETTE["green_dark"])
    draw_box(draw, (1480, 285, 1655, 425), "ML", "Normal / abnormal per mode", PALETTE["green"], PALETTE["green_dark"])
    arrow(draw, (1085, 355), (1165, 355))
    arrow(draw, (1400, 355), (1480, 355))

    draw_swimlane(draw, (805, 565, 1690, 900), "PDS 3 - Edge AI / TinyML", PALETTE["orange_dark"])
    draw_box(draw, (850, 675, 1085, 815), "Codec / Splitter", "Pisahkan audio dengar dan AI", PALETTE["gray"], PALETTE["line"])
    draw_box(draw, (1165, 675, 1400, 815), "MCU TinyML", "Mode-specific features + CNN", PALETTE["orange"], PALETTE["orange_dark"])
    draw_box(draw, (1480, 675, 1655, 815), "HP / PC", "Dashboard hasil + confidence", PALETTE["green"], PALETTE["green_dark"])
    arrow(draw, (1085, 745), (1165, 745))
    arrow(draw, (1400, 745), (1480, 745))

    draw_split_arrow(draw, (705, 490), [(850, 355), (850, 745)])
    draw_note(
        draw,
        (130, 875, 700, 1035),
        "Low-Cost Multi-Mode",
        "Perangkat punya mode akustik original dan mode digital. Visualisasi tetap diarahkan ke HP/PC agar hardware tetap sederhana.",
    )
    save(img, "01_overall_ai_architecture.png")


def draw_pds2_pipeline() -> None:
    img, draw = canvas("PDS 2 AI Pipeline", "Inference multi-mode dilakukan di smartphone atau PC")
    draw_swimlane(draw, (55, 230, 760, 650), "Embedded Acquisition Layer", PALETTE["blue_dark"])
    draw_swimlane(draw, (800, 230, 1745, 650), "Application / AI Layer", PALETTE["green_dark"])
    boxes = [
        ((95, 365, 285, 535), "Sensor", "Heart / lung / abdomen"),
        ((345, 365, 535, 535), "MCU", "Gain, filter, window"),
        ((595, 365, 745, 535), "WS", "WebSocket\nWi-Fi stream"),
        ((830, 365, 1010, 535), "Receive", "Buffer audio / fitur"),
        ((1070, 365, 1250, 535), "Mode", "Heart / lung / abdomen"),
        ((1310, 365, 1490, 535), "Features", "MFCC / log-mel"),
        ((1550, 365, 1725, 535), "ML Model", "SVM per mode"),
    ]
    colors = [PALETTE["blue"], PALETTE["blue"], PALETTE["gray"], PALETTE["gray"], PALETTE["green"], PALETTE["green"], PALETTE["green"]]
    outlines = [PALETTE["blue_dark"], PALETTE["blue_dark"], PALETTE["line"], PALETTE["line"], PALETTE["green_dark"], PALETTE["green_dark"], PALETTE["green_dark"]]
    for (xy, title, body), fill, outline in zip(boxes, colors, outlines):
        draw_box(draw, xy, title, body, fill, outline)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][0][2], 450), (boxes[i + 1][0][0], 450))
    draw_box(draw, (710, 780, 1080, 925), "Output Diagnosa", "Normal / Abnormal + confidence", PALETTE["orange"], PALETTE["orange_dark"])
    arrow(draw, (1638, 535), (895, 780))
    draw_note(
        draw,
        (95, 745, 600, 940),
        "Model PDS 2",
        "Gunakan model per mode. SVM direkomendasikan untuk baseline utama; KNN dan Naive Bayes sebagai pembanding.",
    )
    save(img, "02_pds2_smartphone_inference_pipeline.png")


def draw_pds3_pipeline() -> None:
    img, draw = canvas("PDS 3 Edge AI + Audio Jack Pipeline", "Audio jack untuk dokter, AI di mikrokontroler, visualisasi tetap di HP/PC")
    draw_box(draw, (85, 460, 335, 620), "Chestpiece + Mic", "Heart / lung / abdomen sounds", PALETTE["blue"], PALETTE["blue_dark"])
    draw_box(draw, (455, 460, 735, 620), "Codec / Splitter", "Satu input dibagi menjadi dua jalur", PALETTE["gray"], PALETTE["line"])
    arrow(draw, (335, 540), (455, 540))

    draw_swimlane(draw, (840, 190, 1770, 440), "Real-Time Listening Path", PALETTE["green_dark"])
    draw_box(draw, (905, 285, 1215, 405), "Audio Jack 3.5 mm", "Dokter mendengar langsung", PALETTE["green"], PALETTE["green_dark"])
    draw_box(draw, (1335, 285, 1645, 405), "No AI Delay", "Jalur analog tidak menunggu inferensi", PALETTE["green"], PALETTE["green_dark"])
    arrow(draw, (1215, 345), (1335, 345))

    draw_swimlane(draw, (840, 545, 1770, 930), "Edge AI Path", PALETTE["orange_dark"])
    draw_box(draw, (850, 665, 995, 820), "MCU", "ADC / I2S\ncapture", PALETTE["blue"], PALETTE["blue_dark"])
    draw_box(draw, (1030, 665, 1175, 820), "Prep", "Filter +\nsegment", PALETTE["blue"], PALETTE["blue_dark"])
    draw_box(draw, (1210, 665, 1355, 820), "Features", "MFCC /\nlog-mel", PALETTE["orange"], PALETTE["orange_dark"])
    draw_box(draw, (1390, 665, 1535, 820), "TinyML", "Mode CNN\nTFLM int8", PALETTE["orange"], PALETTE["orange_dark"])
    draw_box(draw, (1570, 665, 1740, 820), "HP / PC", "Dashboard\nresult + confidence", PALETTE["green"], PALETTE["green_dark"])
    arrow(draw, (995, 742), (1030, 742))
    arrow(draw, (1175, 742), (1210, 742))
    arrow(draw, (1355, 742), (1390, 742))
    arrow(draw, (1535, 742), (1570, 742))

    draw_split_arrow(draw, (735, 540), [(905, 345), (850, 742)])
    draw_note(
        draw,
        (95, 170, 735, 340),
        "Prinsip Utama",
        "Jalur audio dokter dipisahkan dari jalur AI. Visual tetap dikirim ke HP/PC, tetapi keputusan AI dibuat di mikrokontroler.",
    )
    draw_note(
        draw,
        (95, 720, 735, 925),
        "Model PDS 3",
        "Quantized 1D-CNN dipakai per mode agar inference lokal tetap ringan untuk heart, lung, dan abdomen sounds.",
    )
    save(img, "03_pds3_edge_ai_audio_jack_pipeline.png")


def draw_training_workflow() -> None:
    img, draw = canvas("Machine Learning Workflow", "Alur dari dataset sampai deployment")
    draw_swimlane(draw, (55, 225, 1725, 515), "Training and Evaluation", PALETTE["purple_dark"])
    boxes = [
        ((85, 330, 285, 470), "Dataset", "Heart / lung / abdomen"),
        ((340, 330, 540, 470), "Mode Label", "Normal / abnormal per mode"),
        ((595, 330, 795, 470), "Preprocess", "Resample, denoise, segment"),
        ((850, 330, 1050, 470), "Features", "MFCC / log-mel / delta"),
        ((1105, 330, 1305, 470), "Training", "Mode-specific models"),
        ((1360, 330, 1560, 470), "Evaluation", "Acc, F1,\nCM"),
    ]
    fills = [PALETTE["gray"], PALETTE["gray"], PALETTE["blue"], PALETTE["green"], PALETTE["orange"], PALETTE["purple"]]
    outlines = [PALETTE["line"], PALETTE["line"], PALETTE["blue_dark"], PALETTE["green_dark"], PALETTE["orange_dark"], PALETTE["purple_dark"]]
    for (xy, title, body), fill, outline in zip(boxes, fills, outlines):
        draw_box(draw, xy, title, body, fill, outline)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][0][2], 400), (boxes[i + 1][0][0], 400))

    draw_swimlane(draw, (205, 650, 790, 930), "Deployment Target PDS 2", PALETTE["green_dark"])
    draw_swimlane(draw, (970, 650, 1555, 930), "Deployment Target PDS 3", PALETTE["orange_dark"])
    draw_box(draw, (300, 755, 695, 885), "PDS 2 Export", "Model .pkl dipakai di smartphone / PC", PALETTE["green"], PALETTE["green_dark"])
    draw_box(draw, (1065, 755, 1460, 885), "PDS 3 Export", "Quantize -> .tflite -> C array untuk MCU", PALETTE["orange"], PALETTE["orange_dark"])
    hub_y = 585
    source_x = 1460
    draw.line(((source_x, 470), (source_x, hub_y)), fill=PALETTE["line"], width=4)
    for target_x in (498, 1262):
        draw.line(((source_x, hub_y), (target_x, hub_y), (target_x, 755)), fill=PALETTE["line"], width=4)
        draw.polygon([(target_x, 755), (target_x - 10, 738), (target_x + 10, 738)], fill=PALETTE["line"])
    save(img, "04_ml_training_deployment_workflow.png")


def draw_tinyml_model() -> None:
    img, draw = canvas("TinyML 1D-CNN Model Concept", "Arsitektur ringan per mode untuk mikrokontroler")
    draw_swimlane(draw, (55, 280, 1680, 690), "Inference Model", PALETTE["orange_dark"])
    boxes = [
        ((100, 410, 310, 570), "Input", "Feature window\nper mode"),
        ((395, 410, 605, 570), "Conv1D", "Filter lokal untuk pola waktu"),
        ((690, 410, 900, 570), "Pooling", "Reduksi dimensi dan noise"),
        ((985, 410, 1195, 570), "Dense", "Gabungkan fitur tingkat tinggi"),
        ((1280, 410, 1490, 570), "Softmax", "Probabilitas kelas"),
    ]
    fills = [PALETTE["green"], PALETTE["orange"], PALETTE["orange"], PALETTE["purple"], PALETTE["green"]]
    outlines = [PALETTE["green_dark"], PALETTE["orange_dark"], PALETTE["orange_dark"], PALETTE["purple_dark"], PALETTE["green_dark"]]
    for (xy, title, body), fill, outline in zip(boxes, fills, outlines):
        draw_box(draw, xy, title, body, fill, outline)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][0][2], 490), (boxes[i + 1][0][0], 490))
    draw_box(draw, (1520, 410, 1680, 570), "Output", "Normal /\nAbnormal", PALETTE["green"], PALETTE["green_dark"])
    arrow(draw, (1490, 490), (1520, 490))
    draw_note(
        draw,
        (485, 780, 1305, 930),
        "Deployment Constraint",
        "Gunakan quantization int8, model kecil, dan inference window pendek agar cocok untuk RAM dan flash mikrokontroler.",
    )
    arrow(draw, (895, 690), (895, 780))
    save(img, "05_tinyml_1d_cnn_model_concept.png")


def draw_decision_logic() -> None:
    img, draw = canvas("ML Decision Logic", "Cara model menghasilkan label dan confidence")
    draw_swimlane(draw, (60, 235, 1740, 705), "Per-Window Inference", PALETTE["purple_dark"])
    boxes = [
        ((105, 385, 305, 545), "Audio Window", "Heart / lung / abdomen"),
        ((380, 385, 580, 545), "Preprocess", "Denoise + normalize"),
        ((655, 385, 855, 545), "Features", "MFCC / log-mel"),
        ((930, 385, 1130, 545), "Mode Model", "SVM atau 1D-CNN"),
        ((1205, 385, 1405, 545), "Prob.", "P(Normal), P(Abnormal)"),
    ]
    fills = [PALETTE["gray"], PALETTE["blue"], PALETTE["green"], PALETTE["orange"], PALETTE["purple"]]
    outlines = [PALETTE["line"], PALETTE["blue_dark"], PALETTE["green_dark"], PALETTE["orange_dark"], PALETTE["purple_dark"]]
    for (xy, title, body), fill, outline in zip(boxes, fills, outlines):
        draw_box(draw, xy, title, body, fill, outline)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][0][2], 465), (boxes[i + 1][0][0], 465))

    draw_box(draw, (1495, 305, 1695, 455), "Abnormal", "Jika P(Abnormal) >= threshold", PALETTE["red"], PALETTE["red_dark"])
    draw_box(draw, (1495, 565, 1695, 715), "Normal", "Jika P(Abnormal) < threshold", PALETTE["green"], PALETTE["green_dark"])
    draw_split_arrow(draw, (1405, 465), [(1495, 380), (1495, 640)])
    draw_note(
        draw,
        (305, 790, 1495, 945),
        "Threshold dan Confidence",
        "Threshold disetel per mode dari validation set. Confidence ditampilkan agar pengguna tahu tingkat keyakinan model, bukan hanya label akhir.",
    )
    save(img, "06_ml_decision_logic.png")


def write_mermaid_files() -> None:
    mermaids = {
        "00_auscultation_mode_strategy.mmd": """flowchart LR
    A[Inspired Digital Stethoscope Design] --> B[Acoustic Original Mode]
    A --> C[Digital Mode]
    C --> D1[PDS 2: AI in Smartphone / PC]
    C --> D2[PDS 3: AI in Microcontroller]
    D1 --> E[HP / PC Visual Output]
    D2 --> E
    A --> F[Heart / Lung / Abdomen Modes]
""",
        "01_overall_ai_architecture.mmd": """flowchart LR
    A[Body Sounds: Heart / Lung / Abdomen] --> B[Chestpiece + MEMS Microphone]
    B --> A0[Acoustic Original Mode]
    B --> C1[PDS 2: Microcontroller + Wi-Fi WebSocket]
    C1 --> D1[Smartphone / PC]
    D1 --> E1[Mode-Specific Features + Scaling]
    E1 --> F1[Classical ML per Mode]
    F1 --> G1[Normal / Abnormal + Confidence]
    B --> C2[PDS 3: Codec / Splitter]
    C2 --> D2[Audio Jack 3.5 mm]
    C2 --> E2[Microcontroller]
    E2 --> F2[Mode-Specific Preprocess + Features]
    F2 --> G2[Quantized TinyML Model per Mode]
    G2 --> H2[HP / PC Dashboard Result + Confidence]
""",
        "02_pds2_smartphone_inference_pipeline.mmd": """flowchart LR
    A[Chestpiece + MEMS Mic] --> B[Microcontroller Audio Capture]
    B --> C[Noise Filter / Windowing]
    C --> D[Wi-Fi WebSocket Streaming]
    D --> E[Smartphone / PC]
    E --> F[Mode Selection]
    F --> G[MFCC / Log-Mel Feature Extraction]
    G --> H[SVM / KNN / Naive Bayes per Mode]
    H --> I[Normal / Abnormal]
""",
        "03_pds3_edge_ai_audio_jack_pipeline.mmd": """flowchart LR
    A[Chestpiece + MEMS Mic] --> B[Audio Codec / Splitter]
    B --> C[Audio Jack 3.5 mm]
    B --> D[Microcontroller ADC / I2S]
    D --> E[Preprocessing]
    E --> F[Mode-Specific Features]
    F --> G[Quantized TinyML Model]
    G --> H[Result Stream]
    H --> I[Smartphone / PC Dashboard]
""",
        "04_ml_training_deployment_workflow.mmd": """flowchart LR
    A[Heart / Lung / Abdomen Dataset] --> B[Mode Labels]
    B --> C[Mode-Specific Preprocessing]
    C --> D[MFCC / Log-Mel Features]
    D --> E[Training]
    E --> F[Evaluation]
    F --> G1[PDS 2 Export .pkl]
    F --> G2[PDS 3 Quantized .tflite]
    G2 --> H2[C Array]
    H2 --> I2[Microcontroller Firmware]
""",
        "05_tinyml_1d_cnn_model_concept.mmd": """flowchart LR
    A[Mode-Specific Feature Input] --> B[Conv1D]
    B --> C[Pooling]
    C --> D[Dense]
    D --> E[Softmax]
    E --> F[Normal / Abnormal]
""",
        "06_ml_decision_logic.mmd": """flowchart LR
    A[Heart / Lung / Abdomen Window] --> B[Preprocess]
    B --> C[Mode-Specific Features]
    C --> D[Mode-Specific SVM or Quantized 1D-CNN]
    D --> E[Class Probability]
    E --> F{P(Abnormal) >= threshold?}
    F -->|Yes| G[Abnormal]
    F -->|No| H[Normal]
""",
    }
    for name, text in mermaids.items():
        (DIAGRAM_DIR / name).write_text(text, encoding="utf-8")


def add_heading(document: Document, text: str, level: int) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.color.rgb = RGBColor(29, 37, 44)


def add_paragraph(document: Document, text: str, bold_prefix: str | None = None) -> None:
    p = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text[len(bold_prefix) :])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(8)


def set_run_times_new_roman(run) -> None:
    run.font.name = "Times New Roman"
    if run._element.rPr is not None and run._element.rPr.rFonts is not None:
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")


def apply_document_font(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_times_new_roman(run)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_times_new_roman(run)


def add_picture(document: Document, filename: str, caption: str) -> None:
    document.add_picture(str(DIAGRAM_DIR / filename), width=Inches(6.6))
    last = document.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = document.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def add_diagram_explanation(document: Document, title: str, rows: list[list[str]]) -> None:
    add_heading(document, title, 2)
    add_table(document, ["Bagian Diagram", "Penjelasan Detail"], rows)


def write_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.color.rgb = RGBColor(29, 37, 44)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = subtitle.add_run(SUBTITLE)
    rs.font.size = Pt(11)
    rs.font.name = "Times New Roman"
    rs._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rs.font.color.rgb = RGBColor(88, 97, 106)

    add_heading(document, "Ringkasan Konsep", 1)
    add_paragraph(
        document,
        "Dokumen ini menjelaskan arsitektur AI untuk prototipe stetoskop digital low-cost berbasis mikrokontroler, mikrofon MEMS, dan desain yang terinspirasi dari Littmann digital stethoscope. Perangkat tidak hanya diarahkan untuk suara abdomen/bowel sounds, tetapi juga untuk suara jantung dan paru sebagai mode auskultasi utama.",
    )
    add_paragraph(
        document,
        "PDS 2 digunakan sebagai baseline dengan inference di smartphone atau PC. PDS 3 adalah pengembangan edge AI, yaitu ekstraksi fitur dan inference langsung di mikrokontroler menggunakan TinyML, tetapi tampilan visual hasil tetap diarahkan ke smartphone atau PC. Istilah mikrokontroler dipakai karena pemilihan chip final belum ditetapkan.",
    )
    add_paragraph(
        document,
        "Secara konsep, perangkat memiliki dua cara pemakaian: mode akustik original untuk auskultasi biasa, dan mode digital untuk perekaman, visualisasi HP/PC, serta AI assistive screening. Mode klinisnya tetap multi-mode: Heart Mode untuk phonocardiogram, Lung Mode untuk breath sounds, dan Abdomen Mode untuk bowel sounds. Output AI pada tahap prototipe adalah Normal/Abnormal + confidence per mode, bukan diagnosis klinis final.",
    )

    add_picture(document, "00_auscultation_mode_strategy.png", "Gambar 0. Strategi mode auskultasi jantung, paru, dan abdomen.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 0",
        [
            ["Inspirasi Littmann Digital", "Inspirasi desainnya adalah prinsip stetoskop digital: suara tetap dapat didengar langsung, tetapi juga bisa direkam, divisualisasikan, dan dianalisis secara digital."],
            ["Acoustic Mode", "Mode ini mempertahankan pemakaian seperti stetoskop biasa. Dokter atau tenaga ahli tetap menjadi pengambil keputusan utama melalui pendengaran langsung."],
            ["Digital Mode", "Mode ini mengaktifkan perekaman, streaming, visualisasi HP/PC, dan AI assistive screening. PDS 2 dan PDS 3 berada di cabang digital mode."],
            ["Heart Mode", "Mode ini menangani phonocardiogram, yaitu suara S1/S2 dan kemungkinan murmur. Pipeline AI dapat diarahkan ke klasifikasi Normal/Abnormal atau deteksi indikasi murmur sebagai pengembangan."],
            ["Lung Mode", "Mode ini menangani suara napas seperti vesicular breath sound, wheeze, dan crackle. Karakter paru berbeda dari jantung dan abdomen, sehingga preprocessing dan model perlu mode-specific."],
            ["Abdomen Mode", "Mode ini menangani bowel sounds yang muncul sebagai burst/intermittent sound. Use case awal proyek tetap bisa fokus pada Normal/Abnormal bowel sounds."],
            ["Mode Router", "Aplikasi atau firmware perlu mengetahui lokasi auskultasi agar fitur dan model yang dipakai sesuai dengan mode suara tubuh yang sedang direkam."],
        ],
    )

    add_picture(document, "01_overall_ai_architecture.png", "Gambar 1. Overall AI Architecture PDS 2 dan PDS 3.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 1",
        [
            ["Common Acquisition", "Bagian ini menunjukkan input fisik yang sama untuk dua desain: suara jantung, paru, atau abdomen ditangkap melalui chestpiece/acoustic coupling dan mikrofon MEMS."],
            ["Acoustic Mode", "Jalur ini menunjukkan bahwa perangkat tetap bisa dipakai sebagai stetoskop biasa tanpa menunggu proses digital atau AI."],
            ["PDS 2", "Mikrokontroler menangkap audio dan mengirimkan data melalui Wi-Fi WebSocket. Proses AI berjalan di smartphone atau PC, sehingga desain ini cocok untuk eksperimen awal, streaming data, dan tuning model."],
            ["PDS 3", "Audio dibagi ke jalur real-time listening melalui audio jack dan jalur AI lokal. Mikrokontroler menjalankan preprocessing, feature extraction, dan inference TinyML; HP/PC hanya menerima hasil untuk dashboard visual."],
            ["Vendor-neutral MCU", "Diagram sengaja memakai istilah MCU/mikrokontroler karena chip final belum ditentukan. Syarat minimalnya adalah mendukung input audio, memori cukup, dan komunikasi yang dibutuhkan."],
        ],
    )

    add_heading(document, "Konsep Machine Learning", 1)
    add_paragraph(
        document,
        "Input utama sistem adalah suara tubuh yang bersifat non-stasioner, beramplitudo rendah, dan mudah tercampur noise lingkungan. Suara jantung memiliki pola siklik S1/S2, suara paru memiliki pola napas inspirasi-ekspirasi, sedangkan suara abdomen muncul sebagai bowel burst yang tidak selalu periodik. Karena itu, model tidak langsung menerima audio mentah berdurasi panjang.",
    )
    add_paragraph(
        document,
        "Pendekatan yang direkomendasikan adalah mode-specific pipeline. Heart Mode, Lung Mode, dan Abdomen Mode dapat memakai kerangka yang sama, tetapi parameter filtering, segmentasi, fitur, dan modelnya disesuaikan. MFCC tetap menjadi fitur utama karena ringkas, tetapi log-mel spectrogram, envelope, spectral centroid, zero-crossing rate, dan temporal statistics dapat ditambahkan sesuai mode.",
    )
    add_table(
        document,
        ["Komponen ML", "Pilihan Desain", "Alasan"],
        [
            ["Segmentasi audio", "Window pendek dengan overlap", "Membantu menangkap S1/S2 pada jantung, fase napas pada paru, dan bowel burst pada abdomen."],
            ["Fitur audio", "MFCC/log-mel, envelope, delta, temporal stats", "Fitur disesuaikan per mode agar model tidak memaksakan satu representasi untuk semua suara tubuh."],
            ["Normalisasi fitur", "StandardScaler atau normalisasi mean-variance", "Model seperti SVM sensitif terhadap skala fitur, sehingga fitur harus berada pada rentang yang konsisten."],
            ["Output model", "Probabilitas Normal dan Abnormal per mode", "Probabilitas memudahkan thresholding dan interpretasi tingkat keyakinan model."],
            ["Evaluasi", "Accuracy, precision, recall, F1-score, confusion matrix per mode", "F1-score penting jika data Normal dan Abnormal tidak seimbang."],
        ],
    )
    add_table(
        document,
        ["Mode", "Sinyal Utama", "Fokus AI", "Catatan ML"],
        [
            ["Heart", "S1/S2, rhythm, murmur-like sound", "Normal/Abnormal atau indikasi murmur", "Perlu segmentasi siklus jantung jika dataset memungkinkan."],
            ["Lung", "Breath sounds, wheeze, crackle", "Normal/Abnormal atau deteksi adventitious sounds", "Perlu mempertimbangkan fase inspirasi dan ekspirasi."],
            ["Abdomen", "Bowel sounds, burst, interval, intensity", "Normal/Abnormal bowel activity", "Cocok untuk baseline awal proyek karena target awal sudah jelas."],
        ],
    )

    add_heading(document, "PDS 2: Smartphone / PC Inference", 1)
    add_paragraph(
        document,
        "Pada PDS 2, mikrokontroler berperan sebagai perangkat akuisisi dan komunikasi. Suara jantung, paru, atau abdomen direkam oleh sistem chestpiece + MEMS microphone, diproses ringan untuk mengurangi noise, lalu dikirim melalui Wi-Fi WebSocket ke smartphone atau PC.",
    )
    add_paragraph(
        document,
        "Ekstraksi fitur, scaling, mode selection, dan inference dilakukan di smartphone atau PC. Pendekatan ini cocok untuk validasi awal karena proses eksperimen, debugging, dan evaluasi model lebih mudah dilakukan pada perangkat dengan komputasi lebih besar.",
    )
    add_paragraph(
        document,
        "WebSocket dipilih karena cocok untuk streaming data secara terus-menerus melalui Wi-Fi. Berbeda dari HTTP request biasa yang putus-sambung, WebSocket mempertahankan koneksi dua arah sehingga mikrokontroler dapat mengirim frame audio atau fitur secara kontinu ke aplikasi penerima.",
    )
    add_picture(document, "02_pds2_smartphone_inference_pipeline.png", "Gambar 2. Pipeline AI PDS 2 dengan inference di smartphone atau PC.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 2",
        [
            ["Embedded Acquisition Layer", "Mikrokontroler melakukan akuisisi, gain control sederhana, filtering ringan, dan windowing. Tujuannya menjaga data audio cukup bersih sebelum dikirim."],
            ["WebSocket", "WebSocket digunakan untuk streaming audio window atau fitur awal melalui Wi-Fi. Koneksi dibuat persistent sehingga data dapat dikirim terus-menerus ke aplikasi tanpa request HTTP berulang."],
            ["Application / AI Layer", "Smartphone atau PC menerima data, menyusun buffer, memilih mode auskultasi, mengekstrak fitur, lalu menjalankan model classical ML sesuai mode."],
            ["SVM", "SVM menjadi model utama karena cocok untuk dataset kecil-menengah. KNN dan Naive Bayes digunakan sebagai baseline pembanding."],
            ["Output Diagnosa", "Output tidak hanya label Normal/Abnormal, tetapi juga confidence/probability agar hasil lebih informatif untuk evaluasi."],
        ],
    )
    add_table(
        document,
        ["Komponen", "Peran dalam PDS 2"],
        [
            ["Mikrokontroler", "Audio capture, pre-processing ringan, koneksi Wi-Fi, dan streaming WebSocket."],
            ["Smartphone / PC", "MFCC extraction, inference, visualisasi, dan penyimpanan hasil."],
            ["Fitur audio", "MFCC/log-mel dan fitur temporal yang disesuaikan untuk Heart, Lung, atau Abdomen Mode."],
            ["SVM", "Model utama yang direkomendasikan untuk dataset kecil sampai menengah."],
            ["KNN / Naive Bayes", "Model pembanding untuk baseline sederhana."],
        ],
    )
    add_heading(document, "Detail Model PDS 2", 2)
    add_paragraph(
        document,
        "Model utama yang direkomendasikan untuk PDS 2 adalah SVM berbasis fitur audio per mode. SVM cocok untuk dataset kecil sampai menengah karena mampu membentuk decision boundary yang relatif stabil walaupun jumlah data belum besar. Jika pola fitur tidak linear, kernel RBF dapat dipakai; jika dataset sangat kecil atau ingin model lebih mudah dijelaskan, linear SVM bisa menjadi opsi awal.",
    )
    add_paragraph(
        document,
        "KNN dan Naive Bayes tetap digunakan sebagai baseline pembanding. KNN mudah dipahami dan dapat menunjukkan apakah fitur MFCC sudah memisahkan kelas secara natural. Naive Bayes sangat ringan dan cepat, tetapi asumsi independensi fitur sering kurang cocok untuk audio sehingga lebih tepat sebagai baseline, bukan model final utama.",
    )
    add_table(
        document,
        ["Model", "Peran", "Kelebihan", "Keterbatasan"],
        [
            ["SVM", "Model utama PDS 2", "Stabil untuk dataset kecil-menengah dan bekerja baik pada fitur MFCC yang sudah diskalakan.", "Perlu tuning kernel, C, dan gamma."],
            ["KNN", "Baseline pembanding", "Sederhana dan tidak perlu training kompleks.", "Inference makin berat jika data training banyak."],
            ["Naive Bayes", "Baseline ringan", "Cepat dan mudah dijelaskan.", "Asumsi independensi fitur sering terlalu sederhana untuk sinyal audio."],
        ],
    )

    add_heading(document, "PDS 3: Edge AI / TinyML", 1)
    add_paragraph(
        document,
        "Pada PDS 3, jalur audio dibagi menjadi dua. Jalur pertama masuk ke audio jack 3.5 mm agar dokter atau tenaga ahli dapat mendengar suara secara real-time. Jalur kedua masuk ke mikrokontroler untuk preprocessing, feature extraction, dan mode-specific inference. Tidak ada display onboard khusus; hasil label, confidence, dan metadata dikirim ke smartphone atau PC untuk visualisasi.",
    )
    add_paragraph(
        document,
        "Dengan desain ini, pembeda utama PDS 2 dan PDS 3 bukan lokasi tampilan visualnya, melainkan lokasi komputasi AI. Pada PDS 2, HP/PC menjalankan feature extraction dan model ML. Pada PDS 3, mikrokontroler menjalankan TinyML secara lokal, lalu HP/PC hanya menjadi dashboard, logger, dan antarmuka pengguna.",
    )
    add_paragraph(
        document,
        "Model yang direkomendasikan adalah 1D-CNN ringan yang sudah dikuantisasi, dengan model atau head berbeda untuk Heart, Lung, dan Abdomen Mode. Model ini sesuai untuk pola audio karena dapat mempelajari pola lokal pada domain waktu atau urutan fitur, tetapi tetap harus dibatasi ukuran dan kompleksitasnya agar cocok untuk RAM dan flash mikrokontroler.",
    )
    add_picture(document, "03_pds3_edge_ai_audio_jack_pipeline.png", "Gambar 3. Pipeline PDS 3 dengan audio jack real-time dan inference on-device.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 3",
        [
            ["Codec / Splitter", "Komponen ini membagi sinyal audio menjadi dua jalur. Jalur pertama untuk didengar langsung, jalur kedua untuk analisis AI."],
            ["Real-Time Listening Path", "Jalur audio jack harus independen dari AI agar dokter tetap dapat mendengar suara jantung, paru, atau abdomen tanpa delay inference."],
            ["Edge AI Path", "Mikrokontroler menangkap sinyal audio, melakukan preprocessing, mengekstrak fitur sesuai mode, lalu menjalankan model TinyML secara lokal."],
            ["1D-CNN TinyML", "Model 1D-CNN dipakai karena dapat mengenali pola temporal pada urutan fitur dengan jumlah parameter lebih ringan dibanding RNN/LSTM."],
            ["HP / PC Dashboard", "Smartphone atau PC menampilkan mode auskultasi, label Normal/Abnormal, confidence, riwayat rekaman, dan catatan eksperimen. Data yang dikirim dari mikrokontroler lebih kecil karena berupa hasil inference, bukan seluruh pipeline AI."],
        ],
    )
    add_heading(document, "Detail Model PDS 3", 2)
    add_paragraph(
        document,
        "PDS 3 memakai konsep TinyML: model dilatih di Python/Keras, kemudian dikompresi dan dideploy ke firmware mikrokontroler. Input model berupa feature window per mode. Conv1D digunakan karena dapat menangkap pola lokal antar-frame, misalnya siklus S1/S2 pada jantung, wheeze/crackle pada paru, atau burst interval pada abdomen.",
    )
    add_paragraph(
        document,
        "1D-CNN dipilih dibandingkan RNN/LSTM karena lebih ringan untuk inference embedded. Dibandingkan 2D-CNN spectrogram penuh, 1D-CNN pada MFCC/log-mel window lebih hemat memori dan lebih realistis untuk mikrokontroler low-cost.",
    )
    add_table(
        document,
        ["Komponen TinyML", "Rancangan", "Alasan"],
        [
            ["Input", "Feature window per mode, misalnya MFCC/log-mel", "Lebih ringkas daripada waveform mentah dan cukup informatif untuk klasifikasi audio."],
            ["Model", "Small 1D-CNN per mode", "Menangkap pola temporal lokal dengan jumlah parameter relatif kecil."],
            ["Output", "Softmax atau sigmoid", "Menghasilkan probabilitas kelas Normal dan Abnormal per mode."],
            ["Optimisasi", "Quantization int8", "Mengurangi ukuran model, RAM, dan waktu inference pada mikrokontroler."],
            ["Runtime", "TensorFlow Lite for Microcontrollers", "Dirancang untuk menjalankan model kecil tanpa sistem operasi besar."],
        ],
    )

    add_heading(document, "Workflow Machine Learning", 1)
    add_paragraph(
        document,
        "Dataset dibagi berdasarkan mode auskultasi: Heart, Lung, dan Abdomen. Setiap mode memiliki label Normal/Abnormal dan metadata lokasi perekaman. Audio diproses melalui resampling, filtering, segmentation, dan normalisasi sebelum fitur diekstraksi.",
    )
    add_paragraph(
        document,
        "Untuk PDS 2, model diekspor sebagai file model klasik seperti .pkl. Untuk PDS 3, model Keras dikompresi melalui quantization, dikonversi menjadi .tflite, lalu diubah menjadi C array agar bisa dimasukkan ke firmware mikrokontroler.",
    )
    add_picture(document, "04_ml_training_deployment_workflow.png", "Gambar 4. Workflow training, evaluasi, dan deployment model.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 4",
        [
            ["Dataset", "Audio dikumpulkan per mode: jantung, paru, dan abdomen. Metadata penting meliputi lokasi auskultasi, durasi, sumber data, sampling rate, dan kondisi perekaman."],
            ["Mode Label", "Setiap sampel harus punya mode yang jelas. Tanpa mode label, model bisa mencampur karakter suara yang secara fisiologis berbeda."],
            ["Preprocess", "Audio di-resample, dibersihkan dari noise, dinormalisasi amplitudonya, dan dipotong menjadi window pendek. Parameter filter dapat berbeda untuk jantung, paru, dan abdomen."],
            ["Features", "MFCC, log-mel, delta, envelope, dan temporal statistics dapat digunakan sebagai input. PDS 2 memakai fitur ini untuk SVM; PDS 3 memakai feature window untuk 1D-CNN."],
            ["Deployment", "PDS 2 mengekspor model klasik per mode seperti .pkl. PDS 3 mengekspor model quantized .tflite yang dikonversi ke C array untuk firmware mikrokontroler."],
        ],
    )

    add_heading(document, "Konsep Model TinyML", 1)
    add_paragraph(
        document,
        "Input model PDS 3 berupa feature window berukuran kecil sesuai mode. Blok Conv1D menangkap pola lokal, pooling mengurangi dimensi, dense layer menggabungkan fitur, dan softmax/sigmoid menghasilkan probabilitas kelas Normal atau Abnormal.",
    )
    add_picture(document, "05_tinyml_1d_cnn_model_concept.png", "Gambar 5. Konsep arsitektur 1D-CNN ringan untuk TinyML pada mikrokontroler.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 5",
        [
            ["Input Features", "Input berupa matriks fitur per window audio. Format ini jauh lebih kecil daripada waveform mentah dan lebih cocok untuk mikrokontroler."],
            ["Conv1D", "Conv1D mencari pola lokal pada urutan frame fitur, misalnya S1/S2, wheeze/crackle, bowel burst, transisi energi, atau pola temporal pendek."],
            ["Pooling", "Pooling mengurangi ukuran fitur agar model lebih ringan dan lebih tahan terhadap variasi kecil pada sinyal."],
            ["Dense + Softmax", "Dense layer menggabungkan fitur menjadi keputusan akhir. Softmax/sigmoid menghasilkan probabilitas kelas Normal dan Abnormal per mode."],
            ["Deployment Constraint", "Model harus kecil, dikuantisasi int8, dan memakai window inference pendek supaya sesuai dengan RAM, flash, dan latency mikrokontroler."],
        ],
    )

    add_picture(document, "06_ml_decision_logic.png", "Gambar 6. Decision logic untuk mengubah probabilitas model menjadi label diagnosis.")
    add_diagram_explanation(
        document,
        "Penjelasan Gambar 6",
        [
            ["Audio Window", "Sistem tidak menilai satu rekaman panjang sekaligus, tetapi menilai beberapa potongan pendek sesuai mode auskultasi agar pola penting tidak hilang."],
            ["Probability", "Model menghasilkan probabilitas kelas, misalnya P(Normal) dan P(Abnormal). Nilai ini lebih informatif daripada label tunggal."],
            ["Threshold", "Threshold awal dapat memakai 0.5, tetapi sebaiknya disetel memakai validation set per mode. Threshold lebih tinggi mengurangi false alarm, sedangkan threshold lebih rendah lebih sensitif terhadap abnormal."],
            ["Final Label", "Jika P(Abnormal) melewati threshold mode terkait, sistem menampilkan Abnormal; jika tidak, sistem menampilkan Normal. Confidence ditampilkan sebagai nilai pendukung keputusan."],
        ],
    )

    add_heading(document, "Perbandingan PDS 2 dan PDS 3", 1)
    add_table(
        document,
        ["Aspek", "PDS 2", "PDS 3"],
        [
            ["Lokasi inference", "Smartphone / PC", "Mikrokontroler"],
            ["Output visual", "Smartphone / PC menampilkan hasil AI yang dihitung di aplikasi", "Smartphone / PC menampilkan hasil AI yang dihitung di mikrokontroler"],
            ["Model utama", "SVM berbasis fitur per mode", "Quantized 1D-CNN per mode"],
            ["Kompleksitas embedded", "Rendah", "Tinggi"],
            ["Latency audio dokter", "Bergantung streaming", "Real-time melalui audio jack"],
            ["Kelebihan", "Mudah diuji dan dievaluasi", "AI berjalan lokal sehingga beban HP/PC lebih ringan dan data yang dikirim bisa berupa hasil ringkas"],
            ["Risiko teknis", "Stabilitas Wi-Fi, latency jaringan, buffering, dan packet drop", "RAM, flash, latency inference, optimasi model, dan sinkronisasi hasil ke HP/PC"],
        ],
    )

    add_heading(document, "Rekomendasi Desain", 1)
    add_paragraph(
        document,
        "Rekomendasi implementasi bertahap adalah memvalidasi pipeline PDS 2 terlebih dahulu menggunakan mode-specific features + SVM sebagai baseline. Abdomen Mode dapat menjadi fokus awal karena target bowel sounds sudah jelas, lalu Heart Mode dan Lung Mode ditambahkan sebagai ekspansi dataset dan model.",
    )
    add_paragraph(
        document,
        "Dengan strategi ini, PDS 2 berfungsi sebagai proof-of-concept AI untuk semua mode, sedangkan PDS 3 menjadi prototipe edge-AI yang menambahkan audio jack real-time dan inference lokal, namun tetap memakai HP/PC sebagai tampilan visual agar hardware tetap sederhana dan low-cost.",
    )

    apply_document_font(document)
    primary_docx = OUT_DIR / "AI_Architecture_PDS2_PDS3.docx"
    try:
        document.save(primary_docx)
    except PermissionError:
        fallback_docx = OUT_DIR / "AI_Architecture_PDS2_PDS3_UPDATED.docx"
        document.save(fallback_docx)
        print(f"Primary DOCX is locked; saved fallback DOCX: {fallback_docx}")


def escape_latex(text: str) -> str:
    replacements = {
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def write_latex() -> None:
    tex = r"""\documentclass[12pt,a4paper]{article}
\usepackage[margin=2.5cm]{geometry}
\usepackage{graphicx}
\usepackage{booktabs}
\usepackage{array}
\usepackage{float}
\usepackage{hyperref}
\usepackage{enumitem}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{mathptmx}

\title{Arsitektur AI Stetoskop Digital Multi-Mode}
\author{Proyek PDS 2 dan PDS 3}
\date{\today}

\begin{document}
\maketitle

\section{Ringkasan Konsep}
Dokumen ini menjelaskan arsitektur AI untuk prototipe stetoskop digital low-cost berbasis mikrokontroler, mikrofon MEMS, dan desain yang terinspirasi dari Littmann digital stethoscope. Perangkat tidak hanya diarahkan untuk suara abdomen/bowel sounds, tetapi juga untuk suara jantung dan paru sebagai mode auskultasi utama.

PDS 2 digunakan sebagai baseline dengan inference di smartphone atau PC. PDS 3 adalah pengembangan edge AI, yaitu ekstraksi fitur dan inference langsung di mikrokontroler menggunakan TinyML, tetapi tampilan visual hasil tetap diarahkan ke smartphone atau PC. Secara konsep, perangkat memiliki mode akustik original untuk auskultasi biasa dan mode digital untuk perekaman, visualisasi HP/PC, serta AI assistive screening.

Mode klinisnya tetap multi-mode: Heart Mode untuk phonocardiogram, Lung Mode untuk breath sounds, dan Abdomen Mode untuk bowel sounds. Output AI pada tahap prototipe adalah Normal/Abnormal + confidence per mode, bukan diagnosis klinis final.

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/00_auscultation_mode_strategy.png}
\caption{Strategi mode auskultasi jantung, paru, dan abdomen.}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/01_overall_ai_architecture.png}
\caption{Overall AI Architecture PDS 2 dan PDS 3.}
\end{figure}

\section{Konsep Machine Learning}
Input utama sistem adalah suara tubuh yang bersifat non-stasioner, beramplitudo rendah, dan mudah tercampur noise lingkungan. Suara jantung memiliki pola siklik S1/S2, suara paru memiliki pola napas inspirasi-ekspirasi, sedangkan suara abdomen muncul sebagai bowel burst yang tidak selalu periodik. Karena itu, model tidak langsung menerima audio mentah berdurasi panjang.

Pendekatan yang direkomendasikan adalah mode-specific pipeline. Heart Mode, Lung Mode, dan Abdomen Mode dapat memakai kerangka yang sama, tetapi parameter filtering, segmentasi, fitur, dan modelnya disesuaikan. MFCC tetap menjadi fitur utama karena ringkas, tetapi log-mel spectrogram, envelope, spectral centroid, zero-crossing rate, dan temporal statistics dapat ditambahkan sesuai mode.

\begin{table}[H]
\centering
\begin{tabular}{p{0.25\linewidth}p{0.28\linewidth}p{0.37\linewidth}}
\toprule
\textbf{Komponen ML} & \textbf{Pilihan Desain} & \textbf{Alasan} \\
\midrule
Segmentasi audio & Window pendek dengan overlap & Membantu menangkap S1/S2 pada jantung, fase napas pada paru, dan bowel burst pada abdomen. \\
Fitur audio & MFCC/log-mel, envelope, delta, temporal stats & Fitur disesuaikan per mode agar model tidak memaksakan satu representasi untuk semua suara tubuh. \\
Normalisasi fitur & StandardScaler atau mean-variance normalization & SVM sensitif terhadap skala fitur sehingga fitur harus konsisten. \\
Output model & Probabilitas Normal dan Abnormal & Probabilitas memudahkan thresholding dan interpretasi keyakinan model. \\
Evaluasi & Accuracy, precision, recall, F1-score, confusion matrix per mode & F1-score penting jika data Normal dan Abnormal tidak seimbang. \\
\bottomrule
\end{tabular}
\caption{Konsep machine learning yang digunakan pada sistem.}
\end{table}

\begin{table}[H]
\centering
\begin{tabular}{p{0.17\linewidth}p{0.26\linewidth}p{0.28\linewidth}p{0.21\linewidth}}
\toprule
\textbf{Mode} & \textbf{Sinyal Utama} & \textbf{Fokus AI} & \textbf{Catatan ML} \\
\midrule
Heart & S1/S2, rhythm, murmur-like sound & Normal/Abnormal atau indikasi murmur & Perlu segmentasi siklus jantung jika dataset memungkinkan. \\
Lung & Breath sounds, wheeze, crackle & Normal/Abnormal atau adventitious sounds & Pertimbangkan fase inspirasi dan ekspirasi. \\
Abdomen & Bowel sounds, burst, interval, intensity & Normal/Abnormal bowel activity & Cocok untuk baseline awal proyek. \\
\bottomrule
\end{tabular}
\caption{Mode auskultasi yang didukung oleh desain AI.}
\end{table}

\section{PDS 2: Smartphone / PC Inference}
Pada PDS 2, mikrokontroler berperan sebagai perangkat akuisisi dan komunikasi. Suara jantung, paru, atau abdomen direkam oleh sistem chestpiece + MEMS microphone, diproses ringan untuk mengurangi noise, lalu dikirim melalui Wi-Fi WebSocket ke smartphone atau PC.

Ekstraksi fitur, scaling, mode selection, dan inference dilakukan di smartphone atau PC. Pendekatan ini cocok untuk validasi awal karena proses eksperimen, debugging, dan evaluasi model lebih mudah dilakukan pada perangkat dengan komputasi lebih besar.

WebSocket dipilih karena cocok untuk streaming data secara terus-menerus melalui Wi-Fi. Berbeda dari HTTP request biasa yang putus-sambung, WebSocket mempertahankan koneksi dua arah sehingga mikrokontroler dapat mengirim frame audio atau fitur secara kontinu ke aplikasi penerima.

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/02_pds2_smartphone_inference_pipeline.png}
\caption{Pipeline AI PDS 2 dengan inference di smartphone atau PC.}
\end{figure}

Model utama yang direkomendasikan untuk PDS 2 adalah SVM berbasis fitur audio per mode. KNN dan Naive Bayes tetap berguna sebagai baseline pembanding karena sederhana dan mudah dijelaskan secara akademik.

\subsection{Detail Model PDS 2}
SVM cocok untuk dataset kecil sampai menengah karena mampu membentuk decision boundary yang relatif stabil walaupun jumlah data belum besar. Setiap mode sebaiknya memiliki model atau parameter model sendiri karena distribusi sinyal jantung, paru, dan abdomen berbeda.

KNN dan Naive Bayes tetap digunakan sebagai baseline pembanding. KNN mudah dipahami dan dapat menunjukkan apakah fitur MFCC sudah memisahkan kelas secara natural. Naive Bayes sangat ringan dan cepat, tetapi asumsi independensi fitur sering kurang cocok untuk audio sehingga lebih tepat sebagai baseline, bukan model final utama.

\begin{table}[H]
\centering
\begin{tabular}{p{0.18\linewidth}p{0.23\linewidth}p{0.27\linewidth}p{0.22\linewidth}}
\toprule
\textbf{Model} & \textbf{Peran} & \textbf{Kelebihan} & \textbf{Keterbatasan} \\
\midrule
SVM & Model utama PDS 2 & Stabil untuk dataset kecil-menengah dan bekerja baik pada MFCC yang sudah diskalakan. & Perlu tuning kernel, C, dan gamma. \\
KNN & Baseline pembanding & Sederhana dan tidak perlu training kompleks. & Inference makin berat jika data training banyak. \\
Naive Bayes & Baseline ringan & Cepat dan mudah dijelaskan. & Asumsi independensi fitur sering terlalu sederhana untuk sinyal audio. \\
\bottomrule
\end{tabular}
\caption{Perbandingan model classical machine learning untuk PDS 2.}
\end{table}

\section{PDS 3: Edge AI / TinyML}
Pada PDS 3, jalur audio dibagi menjadi dua. Jalur pertama masuk ke audio jack 3.5 mm agar dokter atau tenaga ahli dapat mendengar suara secara real-time. Jalur kedua masuk ke mikrokontroler untuk preprocessing, feature extraction, dan mode-specific inference. Tidak ada display onboard khusus; hasil label, confidence, dan metadata dikirim ke smartphone atau PC untuk visualisasi.

Dengan desain ini, pembeda utama PDS 2 dan PDS 3 bukan lokasi tampilan visualnya, melainkan lokasi komputasi AI. Pada PDS 2, HP/PC menjalankan feature extraction dan model ML. Pada PDS 3, mikrokontroler menjalankan TinyML secara lokal, lalu HP/PC hanya menjadi dashboard, logger, dan antarmuka pengguna.

Model yang direkomendasikan adalah 1D-CNN ringan yang sudah dikuantisasi, dengan model atau head berbeda untuk Heart, Lung, dan Abdomen Mode. Model ini sesuai untuk pola audio karena dapat mempelajari pola lokal pada domain waktu atau urutan fitur, tetapi tetap harus dibatasi ukuran dan kompleksitasnya agar cocok untuk RAM dan flash mikrokontroler.

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/03_pds3_edge_ai_audio_jack_pipeline.png}
\caption{Pipeline PDS 3 dengan audio jack real-time dan inference on-device.}
\end{figure}

\subsection{Detail Model PDS 3}
PDS 3 memakai konsep TinyML: model dilatih di Python/Keras, kemudian dikompresi dan dideploy ke firmware mikrokontroler. Input model berupa feature window per mode. Conv1D digunakan karena dapat menangkap pola lokal antar-frame, misalnya siklus S1/S2 pada jantung, wheeze/crackle pada paru, atau burst interval pada abdomen.

1D-CNN dipilih dibandingkan RNN/LSTM karena lebih ringan untuk inference embedded. Dibandingkan 2D-CNN spectrogram penuh, 1D-CNN pada MFCC/log-mel window lebih hemat memori dan lebih realistis untuk mikrokontroler low-cost.

\begin{table}[H]
\centering
\begin{tabular}{p{0.23\linewidth}p{0.32\linewidth}p{0.35\linewidth}}
\toprule
\textbf{Komponen TinyML} & \textbf{Rancangan} & \textbf{Alasan} \\
\midrule
Input & MFCC window, misalnya 20--40 koefisien per frame & Lebih ringkas daripada waveform mentah dan cukup informatif untuk klasifikasi audio. \\
Model & Small 1D-CNN & Menangkap pola temporal lokal dengan jumlah parameter relatif kecil. \\
Output & Softmax atau sigmoid & Menghasilkan probabilitas kelas Normal dan Abnormal. \\
Optimisasi & Quantization int8 & Mengurangi ukuran model, RAM, dan waktu inference pada mikrokontroler. \\
Runtime & TensorFlow Lite for Microcontrollers & Dirancang untuk menjalankan model kecil tanpa sistem operasi besar. \\
\bottomrule
\end{tabular}
\caption{Rancangan model TinyML untuk PDS 3.}
\end{table}

\section{Workflow Machine Learning}
Dataset dibagi berdasarkan mode auskultasi: Heart, Lung, dan Abdomen. Setiap mode memiliki label Normal/Abnormal dan metadata lokasi perekaman. Audio diproses melalui resampling, filtering, segmentation, dan normalisasi sebelum fitur diekstraksi.

Untuk PDS 2, model diekspor sebagai file model klasik seperti \texttt{.pkl}. Untuk PDS 3, model Keras dikompresi melalui quantization, dikonversi menjadi \texttt{.tflite}, lalu diubah menjadi C array agar bisa dimasukkan ke firmware mikrokontroler.

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/04_ml_training_deployment_workflow.png}
\caption{Workflow training, evaluasi, dan deployment model.}
\end{figure}

\section{Konsep Model TinyML}
Input model PDS 3 berupa feature window berukuran kecil sesuai mode. Blok Conv1D menangkap pola lokal, pooling mengurangi dimensi, dense layer menggabungkan fitur, dan softmax/sigmoid menghasilkan probabilitas kelas Normal atau Abnormal.

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/05_tinyml_1d_cnn_model_concept.png}
\caption{Konsep arsitektur 1D-CNN ringan untuk TinyML pada mikrokontroler.}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\linewidth]{diagrams/06_ml_decision_logic.png}
\caption{Decision logic untuk mengubah probabilitas model menjadi label diagnosis.}
\end{figure}

\section{Penjelasan Detail Diagram}
\subsection*{Gambar 1: Overall AI Architecture}
Common acquisition menunjukkan input fisik yang sama untuk PDS 2 dan PDS 3, yaitu suara jantung, paru, atau abdomen yang ditangkap melalui chestpiece/acoustic coupling dan MEMS microphone. Perangkat tetap memiliki mode akustik original untuk auskultasi biasa. PDS 2 menempatkan AI di smartphone atau PC, sedangkan PDS 3 menempatkan inference di mikrokontroler dan mengirim hasil ke HP/PC untuk visualisasi.

\subsection*{Gambar 2: PDS 2 AI Pipeline}
PDS 2 memisahkan embedded acquisition layer dan application/AI layer. Mikrokontroler menangkap audio, melakukan filtering ringan, dan streaming data melalui Wi-Fi WebSocket. Smartphone atau PC menjalankan buffer, mode selection, feature extraction, dan SVM per mode.

\subsection*{Gambar 3: PDS 3 Edge AI + Audio Jack}
PDS 3 memiliki dua jalur audio. Jalur listening path langsung menuju audio jack agar dokter mendengar suara tanpa delay inference. Jalur edge AI path menuju mikrokontroler untuk preprocessing, fitur per mode, dan quantized TinyML model. Hasil inference dikirim ke smartphone atau PC sebagai dashboard visual, bukan ke display onboard khusus. Pemisahan jalur ini penting karena AI tidak boleh mengganggu fungsi utama stetoskop sebagai alat auskultasi.

\subsection*{Gambar 4: Machine Learning Workflow}
Workflow ML dimulai dari dataset heart/lung/abdomen berlabel Normal dan Abnormal, dilanjutkan mode label, preprocessing, feature extraction, training, evaluation, dan deployment. PDS 2 mengekspor model klasik seperti \texttt{.pkl}. PDS 3 mengekspor model \texttt{.tflite} yang dikuantisasi dan dikonversi menjadi C array untuk firmware mikrokontroler.

\subsection*{Gambar 5: TinyML 1D-CNN Model}
Input model adalah feature window per mode. Conv1D menangkap pola temporal lokal, pooling mengurangi dimensi, dense layer menggabungkan fitur, dan softmax/sigmoid menghasilkan probabilitas kelas. Constraint utama adalah ukuran model, RAM, flash, dan latency inference.

\subsection*{Gambar 6: ML Decision Logic}
Model menghasilkan probabilitas kelas, bukan hanya label. Jika probabilitas Abnormal melewati threshold, sistem menampilkan Abnormal; jika tidak, sistem menampilkan Normal. Threshold awal dapat memakai 0.5, tetapi nilai final sebaiknya disetel memakai validation set agar trade-off false alarm dan missed abnormal lebih terkontrol.

\section{Perbandingan PDS 2 dan PDS 3}
\begin{table}[H]
\centering
\begin{tabular}{p{0.25\linewidth}p{0.33\linewidth}p{0.33\linewidth}}
\toprule
\textbf{Aspek} & \textbf{PDS 2} & \textbf{PDS 3} \\
\midrule
Lokasi inference & Smartphone / PC & Mikrokontroler \\
Output visual & Smartphone / PC menampilkan hasil AI yang dihitung di aplikasi & Smartphone / PC menampilkan hasil AI yang dihitung di mikrokontroler \\
Model utama & SVM berbasis fitur per mode & Quantized 1D-CNN per mode \\
Kompleksitas embedded & Rendah & Tinggi \\
Latency audio dokter & Bergantung streaming & Real-time melalui audio jack \\
Kelebihan & Mudah diuji dan dievaluasi & AI berjalan lokal sehingga beban HP/PC lebih ringan dan data yang dikirim bisa berupa hasil ringkas \\
Risiko teknis & Stabilitas Wi-Fi, latency jaringan, buffering, dan packet drop & RAM, flash, latency inference, optimasi model, dan sinkronisasi hasil ke HP/PC \\
\bottomrule
\end{tabular}
\caption{Perbandingan desain AI PDS 2 dan PDS 3.}
\end{table}

\section{Rekomendasi Desain}
Rekomendasi implementasi bertahap adalah memvalidasi pipeline PDS 2 terlebih dahulu menggunakan mode-specific features + SVM sebagai baseline. Abdomen Mode dapat menjadi fokus awal karena target bowel sounds sudah jelas, lalu Heart Mode dan Lung Mode ditambahkan sebagai ekspansi dataset dan model.

Dengan strategi ini, PDS 2 berfungsi sebagai proof-of-concept AI untuk semua mode, sedangkan PDS 3 menjadi prototipe edge-AI yang menambahkan audio jack real-time dan inference lokal, namun tetap memakai HP/PC sebagai tampilan visual agar hardware tetap sederhana dan low-cost.

\end{document}
"""
    (OUT_DIR / "AI_Architecture_PDS2_PDS3.tex").write_text(tex, encoding="utf-8")


def write_readme() -> None:
    readme = """# AI Architecture Stetoskop Digital Multi-Mode

Folder ini berisi dokumen konsep AI untuk prototipe stetoskop digital multi-mode: heart sounds, lung sounds, dan abdomen/bowel sounds.

## File utama

- `AI_Architecture_PDS2_PDS3.docx`: dokumen Word dengan diagram dan penjelasan.
- `AI_Architecture_PDS2_PDS3.tex`: versi LaTeX untuk laporan akademik.
- `diagrams/*.png`: diagram siap pakai untuk Word/PowerPoint.
- `diagrams/*.mmd`: diagram Mermaid yang bisa diedit ulang.

## Diagram

- `00_auscultation_mode_strategy`: strategi mode auskultasi jantung, paru, dan abdomen.
- `01_overall_ai_architecture`: perbandingan arsitektur PDS 2 dan PDS 3.
- `02_pds2_smartphone_inference_pipeline`: pipeline PDS 2.
- `03_pds3_edge_ai_audio_jack_pipeline`: pipeline PDS 3.
- `04_ml_training_deployment_workflow`: workflow training dan deployment.
- `05_tinyml_1d_cnn_model_concept`: konsep model 1D-CNN TinyML.
- `06_ml_decision_logic`: logika probabilitas, threshold, dan label akhir.
"""
    (OUT_DIR / "README.md").write_text(readme, encoding="utf-8")


def main() -> None:
    ensure_dirs()
    draw_auscultation_modes()
    draw_overall_architecture()
    draw_pds2_pipeline()
    draw_pds3_pipeline()
    draw_training_workflow()
    draw_tinyml_model()
    draw_decision_logic()
    write_mermaid_files()
    write_docx()
    write_latex()
    write_readme()
    print(f"Generated files in: {OUT_DIR}")


if __name__ == "__main__":
    main()
